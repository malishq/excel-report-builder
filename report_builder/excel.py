from __future__ import annotations

import math
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .cleaning import CleanStats
from .config import FIELDS, LABELS, Config
from .formatting import fmt_date, month_short, plural
from .metrics import DIMENSIONS, OTHER_LABELS, Report
from .readers import SourceFile

ACCENT = "2A78D6"
ACCENT_DARK = "1C5CAB"
ACCENT_LIGHT = "86B6EF"
SURFACE = "FFFFFF"
INK = "0B0B0B"
INK_SECONDARY = "52514E"
MUTED = "898781"
GRID = "E1E0D9"
BASELINE = "C3C2B7"
OTHER_COLOR = "C3C2B7"
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
DROP_FILL = "FFC7CE"
DROP_FONT = "9C0006"
RISE_FONT = "006300"

SHEET_SUMMARY = "Сводка"
SHEET_MONTHS = "По месяцам"
SHEET_DAYS = "По дням"
SHEET_DATA = "Данные"
SHEET_NAMES = {"manager": "По менеджерам", "category": "По категориям", "region": "По регионам"}
CHART_TITLES = {"manager": "Выручка по менеджерам", "category": "Доля категорий в выручке", "region": "Выручка по регионам"}
TOP_TITLES = {"manager": "Топ-{n} менеджеров", "category": "Топ-{n} категорий", "region": "Топ-{n} регионов"}


@dataclass
class ColumnSpec:
    key: str
    header: str
    number_format: Optional[str] = None
    width: float = 14


@dataclass
class TableRange:
    ws: Worksheet
    header_row: int
    last_row: int
    columns: Dict[str, int] = field(default_factory=dict)

    @property
    def first_row(self) -> int:
        return self.header_row + 1

    @property
    def rows(self) -> int:
        return self.last_row - self.header_row

    def column_range(self, key: str, max_row: Optional[int] = None) -> str:
        letter = get_column_letter(self.columns[key])
        return f"{letter}{self.first_row}:{letter}{max_row or self.last_row}"

    def ref(self) -> str:
        first = get_column_letter(min(self.columns.values()))
        last = get_column_letter(max(self.columns.values()))
        return f"{first}{self.header_row}:{last}{self.last_row}"


@dataclass
class BuildInfo:
    path: Path
    sheets: List[str]
    charts: int


class Formats:
    def __init__(self, currency: str):
        suffix = f' "{currency}"' if currency else ""
        self.money = "#,##0" + suffix
        self.money2 = "#,##0.00" + suffix
        self.integer = "#,##0"
        self.pct = "0.0%"
        self.pct_signed = "+0.0%;-0.0%;0.0%"
        self.date = "DD.MM.YYYY"


def cell_value(value):
    if value is None:
        return None
    if isinstance(value, pd.Period):
        return str(value)
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime()
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, np.bool_):
        return bool(value)
    return value


class ExcelReport:
    def __init__(self, report: Report, config: Config, stats: CleanStats, sources: Sequence[SourceFile], built_at: Optional[datetime] = None):
        self.report = report
        self.config = config
        self.stats = stats
        self.sources = list(sources)
        self.built_at = built_at or datetime.now()
        self.fmt = Formats(config.currency)
        self.wb = Workbook()
        self.charts = 0
        self.header_font = Font(bold=True, color=SURFACE)
        self.header_fill = PatternFill("solid", fgColor=ACCENT_DARK)
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.total_border = Border(top=Side(style="thin", color=BASELINE))

    def save(self, out_path: Path) -> BuildInfo:
        summary = self.wb.active
        summary.title = SHEET_SUMMARY
        months = self.sheet_months()
        dims = {key: self.sheet_dimension(key) for key in DIMENSIONS if key in self.report.by_dim}
        days = self.sheet_days()
        self.sheet_data()
        top_tables = self.sheet_summary(summary, months, dims, days)
        if "category" in dims and "category" in top_tables:
            labels = list(self.report.top["category"]["name"])
            chart = self.pie_chart(top_tables["category"], labels, CHART_TITLES["category"])
            dims["category"].ws.add_chart(chart, self.anchor_after(dims["category"]))
        self.wb.properties.title = self.config.title
        self.wb.properties.creator = "report_builder"
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out_path)
        return BuildInfo(path=out_path, sheets=self.wb.sheetnames, charts=self.charts)

    def set_width(self, ws: Worksheet, column: int, width: float) -> None:
        letter = get_column_letter(column)
        dimension = ws.column_dimensions[letter]
        current = dimension.width if dimension.customWidth else 0
        dimension.width = max(current, width)

    def write_table(self, ws: Worksheet, frame: pd.DataFrame, specs: Sequence[ColumnSpec], start_row: int = 1, start_col: int = 1, totals: Optional[Dict[str, object]] = None) -> TableRange:
        specs = [spec for spec in specs if spec.key in frame.columns]
        table = TableRange(ws=ws, header_row=start_row, last_row=start_row + len(frame))
        for offset, spec in enumerate(specs):
            column = start_col + offset
            cell = ws.cell(row=start_row, column=column, value=spec.header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            table.columns[spec.key] = column
            self.set_width(ws, column, spec.width)
        ws.row_dimensions[start_row].height = 30
        keys = [spec.key for spec in specs]
        for row_offset, values in enumerate(frame[keys].itertuples(index=False), start=1):
            row = start_row + row_offset
            for offset, (spec, value) in enumerate(zip(specs, values)):
                cell = ws.cell(row=row, column=start_col + offset, value=cell_value(value))
                if spec.number_format:
                    cell.number_format = spec.number_format
        if totals:
            row = table.last_row + 1
            for offset, spec in enumerate(specs):
                value = totals.get(spec.key)
                cell = ws.cell(row=row, column=start_col + offset, value=cell_value(value))
                cell.font = Font(bold=True)
                cell.border = self.total_border
                if spec.number_format and value is not None and not isinstance(value, str):
                    cell.number_format = spec.number_format
        return table

    def finish_sheet(self, ws: Worksheet, table: TableRange) -> None:
        ws.freeze_panes = ws.cell(row=table.first_row, column=1)
        ws.auto_filter.ref = table.ref()
        ws.sheet_view.zoomScale = 110

    def color_scale(self, table: TableRange, key: str) -> None:
        if key in table.columns and table.rows > 1:
            rule = ColorScaleRule(start_type="min", start_color=SURFACE, end_type="max", end_color=ACCENT_LIGHT)
            table.ws.conditional_formatting.add(table.column_range(key), rule)

    def highlight_growth(self, table: TableRange, key: str) -> None:
        if key not in table.columns:
            return
        rng = table.column_range(key)
        drop = CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill(bgColor=DROP_FILL), font=Font(color=DROP_FONT))
        rise = CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=RISE_FONT))
        table.ws.conditional_formatting.add(rng, drop)
        table.ws.conditional_formatting.add(rng, rise)

    def data_bar(self, table: TableRange, key: str) -> None:
        if key in table.columns:
            rule = DataBarRule(start_type="num", start_value=0, end_type="max", color=ACCENT_LIGHT)
            table.ws.conditional_formatting.add(table.column_range(key), rule)

    def style_chart(self, chart) -> None:
        chart.legend = None
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=GRID)))
        chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=BASELINE))
        self.charts += 1

    def bar_chart(self, table: TableRange, cat_key: str, val_key: str, title: str, horizontal: bool = False, max_rows: Optional[int] = None) -> BarChart:
        last_row = table.last_row if max_rows is None else min(table.last_row, table.header_row + max_rows)
        cats = Reference(table.ws, min_col=table.columns[cat_key], min_row=table.first_row, max_row=last_row)
        vals = Reference(table.ws, min_col=table.columns[val_key], min_row=table.header_row, max_row=last_row)
        chart = BarChart()
        chart.type = "bar" if horizontal else "col"
        chart.title = title
        chart.add_data(vals, titles_from_data=True)
        chart.set_categories(cats)
        self.style_chart(chart)
        series = chart.series[0]
        series.graphicalProperties.solidFill = ACCENT
        series.graphicalProperties.line.noFill = True
        chart.gapWidth = 70
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.numFmt = "#,##0"
        if horizontal:
            chart.x_axis.scaling.orientation = "maxMin"
            chart.y_axis.crosses = "max"
        chart.width = 17
        chart.height = 8.5
        return chart

    def pie_chart(self, table: TableRange, labels: Sequence[str], title: str) -> PieChart:
        cats = Reference(table.ws, min_col=table.columns["name"], min_row=table.first_row, max_row=table.last_row)
        vals = Reference(table.ws, min_col=table.columns["revenue"], min_row=table.header_row, max_row=table.last_row)
        chart = PieChart()
        chart.title = title
        chart.add_data(vals, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showSerName = False
        chart.dataLabels.showLeaderLines = True
        series = chart.series[0]
        for index, label in enumerate(labels):
            point = DataPoint(idx=index)
            color = OTHER_COLOR if label in OTHER_LABELS else PALETTE[index % len(PALETTE)]
            point.graphicalProperties.solidFill = color
            point.graphicalProperties.line.solidFill = SURFACE
            point.graphicalProperties.line.width = 19050
            series.dPt.append(point)
        chart.legend.position = "r"
        chart.width = 17
        chart.height = 8.5
        self.charts += 1
        return chart

    def line_chart(self, table: TableRange, title: str) -> LineChart:
        cats = Reference(table.ws, min_col=table.columns["date"], min_row=table.first_row, max_row=table.last_row)
        vals = Reference(table.ws, min_col=table.columns["revenue"], min_row=table.header_row, max_row=table.last_row)
        chart = LineChart()
        chart.title = title
        chart.add_data(vals, titles_from_data=True)
        chart.set_categories(cats)
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.number_format = "DD.MM"
        chart.x_axis.majorTimeUnit = "days"
        chart.x_axis.majorUnit = 7
        chart.y_axis.crossAx = 500
        self.style_chart(chart)
        series = chart.series[0]
        series.graphicalProperties.line.solidFill = ACCENT
        series.graphicalProperties.line.width = 19050
        series.smooth = False
        series.marker.symbol = "none"
        chart.width = 34
        chart.height = 9
        return chart

    def anchor_after(self, table: TableRange, row: int = 2) -> str:
        column = max(table.columns.values()) + 2
        return f"{get_column_letter(column)}{row}"

    def sheet_months(self) -> TableRange:
        ws = self.wb.create_sheet(SHEET_MONTHS)
        report = self.report
        specs = [
            ColumnSpec("label", "Месяц", None, 18),
            ColumnSpec("revenue", "Выручка", self.fmt.money, 18),
            ColumnSpec("orders", "Заказы", self.fmt.integer, 11),
            ColumnSpec("quantity", "Продано, шт", self.fmt.integer, 14),
            ColumnSpec("avg_check", "Средний чек", self.fmt.money, 16),
            ColumnSpec("growth", "Рост к пред. месяцу", self.fmt.pct_signed, 20),
        ]
        totals = {"label": "Итого", "revenue": report.revenue, "orders": report.orders, "quantity": report.quantity, "avg_check": report.avg_check}
        table = self.write_table(ws, report.by_month, specs, totals=totals)
        self.finish_sheet(ws, table)
        self.color_scale(table, "revenue")
        self.highlight_growth(table, "growth")
        ws.add_chart(self.bar_chart(table, "label", "revenue", "Выручка по месяцам"), self.anchor_after(table))
        return table

    def sheet_dimension(self, key: str) -> TableRange:
        ws = self.wb.create_sheet(SHEET_NAMES[key])
        report = self.report
        frame = report.by_dim[key]
        last_label = month_short(report.months[-1]) if len(report.months) >= 2 else ""
        prev_label = month_short(report.months[-2]) if len(report.months) >= 2 else ""
        specs = [
            ColumnSpec("name", LABELS[key], None, 28),
            ColumnSpec("revenue", "Выручка", self.fmt.money, 18),
            ColumnSpec("orders", "Заказы", self.fmt.integer, 11),
            ColumnSpec("quantity", "Продано, шт", self.fmt.integer, 14),
            ColumnSpec("avg_check", "Средний чек", self.fmt.money, 16),
            ColumnSpec("share", "Доля в выручке", self.fmt.pct, 15),
            ColumnSpec("prev_month", f"Выручка, {prev_label}", self.fmt.money, 20),
            ColumnSpec("last_month", f"Выручка, {last_label}", self.fmt.money, 20),
            ColumnSpec("growth", "Рост к пред. месяцу", self.fmt.pct_signed, 20),
        ]
        totals = {
            "name": "Итого",
            "revenue": report.revenue,
            "orders": report.orders,
            "quantity": report.quantity,
            "avg_check": report.avg_check,
            "share": 1.0 if report.revenue else None,
        }
        if "last_month" in frame.columns:
            totals["last_month"] = float(frame["last_month"].sum())
            totals["prev_month"] = float(frame["prev_month"].sum())
            totals["growth"] = report.growth
        table = self.write_table(ws, frame, specs, totals=totals)
        self.finish_sheet(ws, table)
        self.color_scale(table, "revenue")
        self.data_bar(table, "share")
        self.highlight_growth(table, "growth")
        if key != "category":
            chart = self.bar_chart(table, "name", "revenue", CHART_TITLES[key], horizontal=True, max_rows=10)
            ws.add_chart(chart, self.anchor_after(table))
        return table

    def sheet_days(self) -> TableRange:
        ws = self.wb.create_sheet(SHEET_DAYS)
        specs = [
            ColumnSpec("date", "Дата", self.fmt.date, 13),
            ColumnSpec("revenue", "Выручка", self.fmt.money, 18),
            ColumnSpec("orders", "Заказы", self.fmt.integer, 11),
            ColumnSpec("avg_check", "Средний чек", self.fmt.money, 16),
        ]
        table = self.write_table(ws, self.report.by_day, specs)
        self.finish_sheet(ws, table)
        self.color_scale(table, "revenue")
        ws.add_chart(self.line_chart(table, "Выручка по дням"), self.anchor_after(table))
        return table

    def sheet_data(self) -> None:
        ws = self.wb.create_sheet(SHEET_DATA)
        data = self.report.data
        keys = [key for key in FIELDS if key in data.columns] + ["source"]
        headers = [LABELS.get(key, "Файл") for key in keys]
        formats = {"date": self.fmt.date, "amount": self.fmt.money2, "quantity": self.fmt.integer}
        widths = {"date": 13, "amount": 16, "quantity": 12, "source": 24}
        ws.append(headers)
        for values in data[keys].itertuples(index=False):
            ws.append([cell_value(value) for value in values])
        last_row = len(data) + 1
        for index, key in enumerate(keys, start=1):
            number_format = formats.get(key)
            if number_format:
                for (cell,) in ws.iter_rows(min_row=2, max_row=last_row, min_col=index, max_col=index):
                    cell.number_format = number_format
            if key in widths:
                width = widths[key]
            else:
                longest = int(data[key].astype(str).str.len().max()) if len(data) else 0
                width = min(max(longest, len(headers[index - 1])) + 3, 45)
            self.set_width(ws, index, width)
        ref = f"A1:{get_column_letter(len(keys))}{last_row}"
        table = Table(displayName="Data", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "A2"

    def sheet_summary(self, ws: Worksheet, months: TableRange, dims: Dict[str, TableRange], days: TableRange) -> Dict[str, TableRange]:
        report = self.report
        config = self.config
        stats = self.stats
        ws.sheet_view.showGridLines = False
        for column, width in ((1, 30), (2, 20), (3, 14), (4, 16), (5, 4)):
            self.set_width(ws, column, width)
        ws["A1"] = config.title
        ws["A1"].font = Font(bold=True, size=18, color=INK)
        ws["A2"] = f"Период: {fmt_date(report.period_start)} — {fmt_date(report.period_end)}"
        ws["A2"].font = Font(size=11, color=INK_SECONDARY)
        ws["A3"] = (
            f"Собрано {self.built_at:%d.%m.%Y %H:%M} из {plural(len(self.sources), 'файла', 'файлов', 'файлов')}, "
            f"{plural(stats.rows_out, 'строка', 'строки', 'строк')} после очистки"
        )
        ws["A3"].font = Font(size=10, color=MUTED)
        row = 5
        ws.cell(row=row, column=1, value="Ключевые показатели").font = Font(bold=True, size=12, color=INK)
        row += 1
        kpis = [
            ("Выручка", report.revenue, self.fmt.money, None),
            ("Заказы", report.orders, self.fmt.integer, None),
            ("Средний чек", report.avg_check, self.fmt.money, None),
        ]
        if report.quantity is not None:
            kpis.append(("Продано, шт", report.quantity, self.fmt.integer, None))
        if report.clients is not None:
            kpis.append(("Клиентов", report.clients, self.fmt.integer, None))
        growth_value = report.growth if report.growth is not None else "—"
        kpis.append(("Рост к прошлому месяцу", growth_value, self.fmt.pct_signed, report.growth_label or None))
        for label, value, number_format, note in kpis:
            ws.cell(row=row, column=1, value=label).font = Font(size=11, color=INK_SECONDARY)
            cell = ws.cell(row=row, column=2, value=cell_value(value))
            cell.font = Font(bold=True, size=12, color=INK)
            cell.alignment = Alignment(horizontal="right")
            if isinstance(value, (int, float)):
                cell.number_format = number_format
            if note:
                ws.cell(row=row, column=3, value=note).font = Font(size=10, color=MUTED)
            if label.startswith("Рост") and isinstance(value, float):
                rng = f"B{row}:B{row}"
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, size=12, color=DROP_FONT)))
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(bold=True, size=12, color=RISE_FONT)))
            row += 1
        top_tables: Dict[str, TableRange] = {}
        for key in ("manager", "category"):
            if key not in report.top:
                continue
            row += 1
            ws.cell(row=row, column=1, value=TOP_TITLES[key].format(n=config.top_n)).font = Font(bold=True, size=12, color=INK)
            row += 1
            specs = [
                ColumnSpec("name", LABELS[key], None, 30),
                ColumnSpec("revenue", "Выручка", self.fmt.money, 20),
                ColumnSpec("orders", "Заказы", self.fmt.integer, 14),
                ColumnSpec("share", "Доля", self.fmt.pct, 16),
            ]
            table = self.write_table(ws, report.top[key], specs, start_row=row)
            self.data_bar(table, "share")
            top_tables[key] = table
            row = table.last_row + 1
        row += 1
        ws.cell(row=row, column=1, value="О данных").font = Font(bold=True, size=12, color=INK)
        row += 1
        for source in self.sources:
            ws.cell(row=row, column=1, value=source.describe()).font = Font(size=10, color=INK_SECONDARY)
            row += 1
        cleaning = (
            f"Исходных строк: {stats.rows_in}; убрано пустых — {stats.empty}, без даты — {stats.bad_date}, "
            f"без суммы — {stats.bad_amount}, дубликатов — {stats.duplicates}"
        )
        if stats.negative:
            cleaning += f", отрицательных — {stats.negative}"
        ws.cell(row=row, column=1, value=cleaning).font = Font(size=10, color=INK_SECONDARY)
        ws.add_chart(self.bar_chart(months, "label", "revenue", "Выручка по месяцам"), "F5")
        if "category" in top_tables:
            labels = list(report.top["category"]["name"])
            ws.add_chart(self.pie_chart(top_tables["category"], labels, CHART_TITLES["category"]), "P5")
        ws.add_chart(self.line_chart(days, "Выручка по дням"), "F23")
        return top_tables


def write_report(report: Report, config: Config, stats: CleanStats, sources: Sequence[SourceFile], out_path: Path, built_at: Optional[datetime] = None) -> BuildInfo:
    return ExcelReport(report, config, stats, sources, built_at).save(Path(out_path))
