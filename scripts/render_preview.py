from __future__ import annotations

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from report_builder.charts_png import setup_style
from report_builder.formatting import fmt_date, fmt_int, fmt_money, fmt_pct

INK = "#0b0b0b"
INK_SECONDARY = "#52514e"
MUTED = "#898781"
HEADER = "#1c5cab"
GRID = "#e1e0d9"
BASELINE = "#c3c2b7"
SCALE_END = "#86b6ef"
DROP_FILL = "#ffc7ce"
DROP_FONT = "#9c0006"
RISE_FONT = "#006300"
STRIPE = "#eef3fa"
TAB_ACTIVE = "#ffffff"
TAB_BAR = "#f3f3f1"
SHEETS = ["Сводка", "По месяцам", "По менеджерам", "По категориям", "По регионам", "По дням", "Данные"]
CAPTION = "Превью отрисовано matplotlib по содержимому examples/report.xlsx (Excel при сборке был недоступен). В Excel оформление, шрифты и диаграммы выглядят иначе."


def format_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    number_format = cell.number_format or ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if "%" in number_format:
            return fmt_pct(value, signed="+" in number_format)
        if '"' in number_format:
            currency = number_format.split('"')[1]
            return fmt_money(value, currency, 2 if ".00" in number_format else 0)
        if "#,##0" in number_format:
            return fmt_int(value)
        return str(value)
    if hasattr(value, "strftime"):
        return fmt_date(value)
    return str(value)


def font_color(cell, default: str = INK) -> str:
    color = cell.font.color
    if color is not None and color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) == 8:
        return "#" + color.rgb[2:].lower()
    return default


def is_header(cell) -> bool:
    fill = cell.fill
    return fill is not None and fill.fill_type == "solid" and str(fill.fgColor.rgb).endswith("1C5CAB")


def blend(start: str, end: str, ratio: float) -> str:
    start_rgb = [int(start[i:i + 2], 16) for i in (1, 3, 5)]
    end_rgb = [int(end[i:i + 2], 16) for i in (1, 3, 5)]
    mixed = [round(a + (b - a) * ratio) for a, b in zip(start_rgb, end_rgb)]
    return "#" + "".join(f"{channel:02x}" for channel in mixed)


class Canvas:
    def __init__(self, width: int, height: int):
        self.width = width
        self.height = height
        self.fig = plt.figure(figsize=(width / 100, height / 100), dpi=100)
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set_xlim(0, width)
        self.ax.set_ylim(height, 0)
        self.ax.axis("off")

    def rect(self, x, y, w, h, color, edge=None, lw=0.8, zorder=1):
        self.ax.add_patch(Rectangle((x, y), w, h, facecolor=color, edgecolor=edge or "none", linewidth=lw, zorder=zorder))

    def line(self, x0, y0, x1, y1, color=GRID, lw=0.8):
        self.ax.plot([x0, x1], [y0, y1], color=color, linewidth=lw, zorder=2, solid_capstyle="butt")

    def text(self, x, y, text, size=11, color=INK, bold=False, ha="left", va="center", zorder=3):
        self.ax.text(x, y, text, fontsize=size, color=color, fontweight="bold" if bold else "normal", ha=ha, va=va, zorder=zorder)

    def image(self, path: Path, x, y, w, h):
        image = plt.imread(path)
        axes = self.fig.add_axes([x / self.width, 1 - (y + h) / self.height, w / self.width, h / self.height])
        axes.imshow(image)
        axes.axis("off")
        self.rect(x, y, w, h, "none", edge=GRID, zorder=4)

    def tabs(self, active: str, y: float):
        self.rect(0, y, self.width, self.height - y, TAB_BAR)
        self.line(0, y, self.width, y, BASELINE)
        x = 24
        for name in SHEETS:
            width = 14 + len(name) * 8.5
            if name == active:
                self.rect(x, y, width, 26, TAB_ACTIVE, edge=BASELINE)
                self.text(x + width / 2, y + 13, name, size=10, color=HEADER, bold=True, ha="center")
            else:
                self.text(x + width / 2, y + 13, name, size=10, color=INK_SECONDARY, ha="center")
            x += width + 4

    def caption(self, y: float):
        self.text(24, y, CAPTION, size=9, color=MUTED)

    def save(self, path: Path):
        self.fig.savefig(path, dpi=100, facecolor="white")
        plt.close(self.fig)


def draw_summary(workbook, screens: Path, out: Path) -> None:
    ws = workbook["Сводка"]
    canvas = Canvas(1600, 980)
    columns = [(40, 330, "left"), (330, 470, "right"), (470, 560, "right"), (560, 640, "right")]
    y = 28.0
    share_max = 0.0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4):
        if row[3].value is not None and isinstance(row[3].value, (int, float)) and row[1].value is not None:
            share_max = max(share_max, float(row[3].value))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4):
        cells = list(row)
        header = is_header(cells[0])
        size = cells[0].font.size or 11
        height = 30 if header else (size * 2.2 if size > 12 else 24)
        if header:
            canvas.rect(columns[0][0] - 8, y - height / 2 + 2, columns[-1][1] - columns[0][0] + 8, height - 4, HEADER)
        elif cells[0].value is not None and cells[1].value is not None and isinstance(cells[3].value, (int, float)):
            width = (columns[3][1] - columns[3][0] - 10) * float(cells[3].value) / share_max if share_max else 0
            canvas.rect(columns[3][0] + 4, y - 8, width, 16, SCALE_END, zorder=1)
            canvas.line(columns[0][0] - 8, y + height / 2 - 2, columns[-1][1], y + height / 2 - 2)
        for (x0, x1, align), cell in zip(columns, cells):
            text = format_cell(cell)
            if not text:
                continue
            if header:
                canvas.text((x0 + x1) / 2, y, text, size=10, color="white", bold=True, ha="center")
                continue
            numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            ha = "right" if numeric or align == "right" and cell.column > 1 and numeric else "left"
            x = x1 - 6 if ha == "right" else x0
            if cell.column == 3 and not numeric:
                x = x0 + 6
            color = font_color(cell)
            if cells[0].value == "Рост к прошлому месяцу" and cell.column == 2 and numeric:
                color = RISE_FONT if cell.value > 0 else DROP_FONT if cell.value < 0 else color
            canvas.text(x, y, text, size=min(size, 15), color=color, bold=bool(cell.font.bold), ha=ha)
        y += height
    canvas.image(screens / "01_revenue_by_month.png", 690, 60, 420, 226)
    canvas.image(screens / "02_categories_share.png", 1140, 60, 420, 226)
    canvas.image(screens / "03_daily_revenue.png", 690, 320, 870, 380)
    canvas.text(690, 735, "Три диаграммы на листе «Сводка» в Excel нативные (столбцы, круг, линия); здесь показаны их PNG-копии из docs/screens.", size=9, color=MUTED)
    canvas.tabs("Сводка", 930)
    canvas.caption(912)
    canvas.save(out)


def draw_managers(workbook, screens: Path, out: Path) -> None:
    ws = workbook["По менеджерам"]
    canvas = Canvas(1600, 620)
    widths = [34, 190, 150, 90, 110, 130, 120, 150, 150, 150]
    x_positions = []
    x = 20
    for width in widths:
        x_positions.append((x, x + width))
        x += width
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9))
    data_rows = [row for row in rows[1:] if row[0].value not in (None, "Итого")]
    revenues = [float(row[1].value) for row in data_rows]
    shares = [float(row[5].value) for row in data_rows]
    low, high = min(revenues), max(revenues)
    y = 44.0
    canvas.rect(x_positions[0][0], y - 16, x_positions[-1][1] - x_positions[0][0], 32, HEADER)
    canvas.text((x_positions[0][0] + x_positions[0][1]) / 2, y, "", size=9)
    for index, cell in enumerate(rows[0]):
        x0, x1 = x_positions[index + 1]
        canvas.text((x0 + x1) / 2, y, str(cell.value), size=9.5, color="white", bold=True, ha="center")
    y += 32
    for number, row in enumerate(rows[1:], start=2):
        total = row[0].value == "Итого"
        height = 28
        if number % 2 == 1 and not total:
            canvas.rect(x_positions[0][0], y - height / 2, x_positions[-1][1] - x_positions[0][0], height, STRIPE, zorder=0)
        if total:
            canvas.line(x_positions[0][0], y - height / 2, x_positions[-1][1], y - height / 2, BASELINE, lw=1.2)
        canvas.text((x_positions[0][0] + x_positions[0][1]) / 2, y, str(number), size=9, color=MUTED, ha="center")
        for index, cell in enumerate(row):
            x0, x1 = x_positions[index + 1]
            text = format_cell(cell)
            numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            color = INK
            if not total and numeric:
                if index == 1 and high > low:
                    canvas.rect(x0, y - height / 2, x1 - x0, height, blend("#ffffff", SCALE_END, (float(cell.value) - low) / (high - low)), zorder=1)
                if index == 5:
                    canvas.rect(x0 + 4, y - 9, (x1 - x0 - 10) * float(cell.value) / max(shares), 18, SCALE_END, zorder=1)
                if index == 8:
                    if cell.value < 0:
                        canvas.rect(x0, y - height / 2, x1 - x0, height, DROP_FILL, zorder=1)
                        color = DROP_FONT
                    elif cell.value > 0:
                        color = RISE_FONT
            if text:
                canvas.text(x1 - 8 if numeric else x0 + 8, y, text, size=10, color=color, bold=total, ha="right" if numeric else "left")
        y += height
    canvas.line(x_positions[0][0], y - 14, x_positions[-1][1], y - 14, GRID)
    canvas.text(20, y + 10, "Выручка: цветовая шкала · Доля: гистограмма в ячейке · Рост: падение подсвечено красным, рост зелёным (правила условного форматирования из файла).", size=9, color=MUTED)
    canvas.image(screens / "04_managers.png", 1300, 30, 280, 150)
    canvas.text(1300, 196, "На листе ещё нативная диаграмма Excel", size=8.5, color=MUTED)
    canvas.tabs("По менеджерам", 570)
    canvas.caption(552)
    canvas.save(out)


def main() -> int:
    report = ROOT / "examples" / "report.xlsx"
    screens = ROOT / "docs" / "screens"
    setup_style()
    workbook = load_workbook(report)
    draw_summary(workbook, screens, screens / "05_summary_preview.png")
    draw_managers(workbook, screens, screens / "06_managers_preview.png")
    print("ok")
    return 0


if __name__ == "__main__":
    sys.exit(main())
