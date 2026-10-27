import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from report_builder.cleaning import prepare
from report_builder.config import Config
from report_builder.excel import write_report
from report_builder.metrics import build_report
from report_builder.readers import read_sources
from report_builder.sample import generate_sample

ROOT = Path(__file__).resolve().parents[1]
EXPECTED_SHEETS = ["Сводка", "По месяцам", "По менеджерам", "По категориям", "По регионам", "По дням", "Данные"]


def build(tmp_path, rows=150):
    data_dir = tmp_path / "data"
    generate_sample(data_dir, seed=7, rows_per_month=rows, start=date(2026, 7, 1))
    config = Config.from_dict(None)
    sources = read_sources([str(data_dir / "*.csv")])
    frame, stats = prepare(sources, config)
    report = build_report(frame, config)
    info = write_report(report, config, stats, sources, tmp_path / "report.xlsx")
    return report, stats, info


def test_sample_files_differ_in_format(tmp_path):
    paths = generate_sample(tmp_path, rows_per_month=100, start=date(2026, 7, 1))
    assert [path.name for path in paths] == ["sales_2026-07.csv", "sales_2026-08.csv", "sales_2026-09.csv"]
    sources = read_sources([str(tmp_path)])
    assert [(source.encoding, source.delimiter) for source in sources] == [("cp1251", ";"), ("utf-8", "\t"), ("utf-8-sig", ",")]
    assert all(source.rows > 50 for source in sources)


def test_workbook_sheets_charts_and_values(tmp_path):
    report, stats, info = build(tmp_path)
    assert info.sheets == EXPECTED_SHEETS
    assert info.charts == 8
    workbook = load_workbook(info.path)
    assert workbook.sheetnames == EXPECTED_SHEETS
    charts = {name: len(workbook[name]._charts) for name in EXPECTED_SHEETS}
    assert charts == {"Сводка": 3, "По месяцам": 1, "По менеджерам": 1, "По категориям": 1, "По регионам": 1, "По дням": 1, "Данные": 0}
    summary = workbook["Сводка"]
    assert summary["A1"].value == "Отчёт по продажам"
    assert summary["A6"].value == "Выручка"
    assert summary["B6"].value == report.revenue
    assert summary["B6"].number_format == '#,##0 "₽"'
    assert summary["B7"].value == report.orders
    assert summary["B11"].value == report.growth
    assert summary["C11"].value == "сен. 2026 к авг. 2026"
    months = workbook["По месяцам"]
    assert [cell.value for cell in months["A"][1:4]] == ["Июль 2026", "Август 2026", "Сентябрь 2026"]
    assert months["A5"].value == "Итого"
    assert months["B5"].value == report.revenue
    assert months.freeze_panes == "A2"
    assert months.auto_filter.ref == "A1:F4"
    assert len(months.conditional_formatting) == 2
    data = workbook["Данные"]
    assert "Data" in data.tables
    assert data.tables["Data"].ref == f"A1:H{stats.rows_out + 1}"
    assert data.freeze_panes == "A2"
    assert [cell.value for cell in data[1]] == ["Дата", "Сумма", "Количество", "Менеджер", "Категория", "Регион", "Клиент", "Файл"]
    assert data["A2"].number_format == "DD.MM.YYYY"
    assert data.max_row == stats.rows_out + 1


def test_workbook_without_optional_columns(tmp_path):
    import pandas as pd
    from report_builder.readers import SourceFile

    source = SourceFile(path=Path("x.csv"), frame=pd.DataFrame({"Дата": ["01.07.2026", "02.08.2026"], "Сумма": ["10", "30"]}))
    config = Config.from_dict(None)
    frame, stats = prepare([source], config)
    report = build_report(frame, config)
    info = write_report(report, config, stats, [source], tmp_path / "small.xlsx")
    assert info.sheets == ["Сводка", "По месяцам", "По дням", "Данные"]
    assert info.charts == 4


def test_cli_end_to_end(tmp_path):
    env = {"PYTHONPATH": str(ROOT)}
    data_dir = tmp_path / "data"
    result = subprocess.run([sys.executable, "-m", "report_builder", "sample", "--out", str(data_dir), "--rows", "120", "--start", "2026-07"], cwd=ROOT, env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    out = tmp_path / "report.xlsx"
    png_dir = tmp_path / "png"
    result = subprocess.run([
        sys.executable, "-m", "report_builder", "build",
        "--input", str(data_dir / "*.csv"),
        "--config", str(ROOT / "config.yml"),
        "--out", str(out),
        "--png", str(png_dir),
    ], cwd=ROOT, env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert out.exists()
    assert "Готово" in result.stdout
    assert sorted(path.name for path in png_dir.glob("*.png")) == ["01_revenue_by_month.png", "02_categories_share.png", "03_daily_revenue.png", "04_managers.png"]


def test_cli_reports_missing_column(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("Дата;Менеджер\n01.07.2026;Анна\n", encoding="utf-8")
    result = subprocess.run([sys.executable, "-m", "report_builder", "build", "--input", str(path), "--out", str(tmp_path / "r.xlsx")], cwd=ROOT, env={"PYTHONPATH": str(ROOT)}, capture_output=True, text=True)
    assert result.returncode == 2
    assert "обязательные колонки" in result.stderr
